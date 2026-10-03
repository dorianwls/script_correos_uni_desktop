from __future__ import annotations

import csv
import html
import json
import queue
import re
import threading
import tempfile
import unicodedata
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from email.utils import getaddresses
from pathlib import Path
from typing import Callable
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook


APP_TITLE = "Verificación de equipos - COMTECH"
APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "datos"
HISTORY_FILE = DATA_DIR / "historial_envios.csv"
OAUTH_CONFIG_FILE = DATA_DIR / "configuracion_oauth.json"
GRAPH_SEND_MAIL_URL = "https://graph.microsoft.com/v1.0/me/sendMail"
GRAPH_SCOPES = ["Mail.Send"]
LEGACY_HEADERS = {
    "numeroserie": "NumeroSerie",
    "marcaymodelo": "Marcaymodelo",
    "caracteristicas": "Caracteristicas",
}
TICKET_HEADERS = {
    "codigo": "Codigo",
    "marca": "Marca",
    "modelo": "Modelo",
    "serie": "Serie",
}
INVENTORY_HEADERS = {
    "items": "Items",
    "noseriecpu": "No. serie CPU",
    "noseriemonitor": "No. serie monitor",
    "responsable": "Responsable",
}
PAIRED_HEADERS = (
    "cpu", "marcaymodelo", "caracteristicasdelequipo",
    "monitores", "marcaymodelo", "caracteristicasdelequipo",
)


@dataclass
class EquipmentRow:
    excel_row: int
    serial: str
    brand_model: str
    characteristics: str
    verification_date: date | None = None
    verification_time: str = ""
    status: str = "Pendiente"
    error: str = ""
    code: str = ""
    brand: str = ""
    model: str = ""
    lab: str = ""
    item: str = ""
    monitor_serial: str = ""
    responsible: str = ""
    inventory_format: bool = False
    paired_format: bool = False
    monitor_brand_model: str = ""
    monitor_characteristics: str = ""

    @property
    def valid(self) -> bool:
        return not self.error

    @property
    def history_identifier(self) -> str:
        if self.paired_format:
            return "cpu-monitor-details:" + json.dumps(
                [self.excel_row, self.serial.casefold(), self.monitor_serial.casefold()],
                ensure_ascii=False,
            )
        if self.inventory_format:
            return "cpu-monitor:" + json.dumps(
                [self.item, self.serial.casefold(), self.monitor_serial.casefold(), self.responsible.casefold()],
                ensure_ascii=False,
            )
        return self.serial.casefold()


@dataclass(frozen=True)
class MailSettings:
    sender: str
    recipient: str
    cc: tuple[str, ...]
    bcc: tuple[str, ...]
    delay_seconds: float
    test_mode: bool


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_date(value: str) -> date:
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError("La fecha debe usar AAAA-MM-DD o DD/MM/AAAA.")


def parse_clock(value: str) -> datetime:
    try:
        return datetime.strptime(value.strip(), "%H:%M")
    except ValueError as exc:
        raise ValueError("Las horas deben usar el formato HH:MM de 24 horas.") from exc


def distribute_uniform_times(start: str, end: str, count: int) -> list[str]:
    """Distribuye horarios dentro de [inicio, fin), excluyendo 12:00–13:00."""
    if count < 0:
        raise ValueError("La cantidad no puede ser negativa.")
    if count == 0:
        return []

    start_dt = parse_clock(start)
    end_dt = parse_clock(end)
    if end_dt <= start_dt:
        raise ValueError("La hora final debe ser posterior a la hora inicial.")

    lunch_start = start_dt.replace(hour=12, minute=0)
    lunch_end = start_dt.replace(hour=13, minute=0)
    segments = [
        (segment_start, segment_end)
        for segment_start, segment_end in (
            (start_dt, min(end_dt, lunch_start)),
            (max(start_dt, lunch_end), end_dt),
        )
        if segment_start < segment_end
    ]
    total_seconds = sum(int((finish - begin).total_seconds()) for begin, finish in segments)
    if total_seconds == 0:
        raise ValueError("El rango está dentro del almuerzo (12:00–13:00). Elija horas fuera de ese intervalo.")
    if count > total_seconds:
        raise ValueError(
            f"El rango solo contiene {total_seconds} segundos fuera del almuerzo y no permite "
            f"asignar {count} horas únicas. Amplíe el rango."
        )

    offsets = [(index * total_seconds) // count for index in range(count)]
    if len(set(offsets)) != count:
        raise ValueError("No fue posible producir horarios únicos con ese rango.")

    show_seconds = total_seconds / count < 60
    fmt = "%I:%M:%S %p" if show_seconds else "%I:%M %p"
    times = []
    for offset in offsets:
        for begin, finish in segments:
            segment_seconds = int((finish - begin).total_seconds())
            if offset < segment_seconds:
                times.append((begin + timedelta(seconds=offset)).strftime(fmt))
                break
            offset -= segment_seconds
    return times


def parse_address_list(value: str, label: str, allow_empty: bool = True) -> tuple[str, ...]:
    value = value.strip()
    if not value:
        if allow_empty:
            return ()
        raise ValueError(f"{label} es obligatorio.")
    parsed = tuple(
        address.strip()
        for _, address in getaddresses([value.replace(";", ",")])
        if address.strip()
    )
    if not parsed:
        raise ValueError(f"{label} no contiene una dirección válida.")
    invalid = [address for address in parsed if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", address)]
    if invalid:
        raise ValueError(f"Dirección inválida en {label}: {invalid[0]}")
    return parsed


def read_inventory_rows(
    raw_headers: Sequence[object], iterator: Iterable[Sequence[object]]
) -> list[EquipmentRow] | None:
    header_map = {normalize_header(header): index for index, header in enumerate(raw_headers)}
    if not any(key in header_map for key in INVENTORY_HEADERS):
        return None
    missing = [label for key, label in INVENTORY_HEADERS.items() if key not in header_map]
    if missing:
        raise ValueError("Faltan columnas: " + ", ".join(missing))
    rows = []
    for excel_row, values in enumerate(iterator, start=2):
        if not any(clean_cell(value) for value in values):
            continue
        fields = {
            key: clean_cell(values[index]) if index < len(values) else ""
            for key, index in header_map.items() if key in INVENTORY_HEADERS
        }
        errors = [f"Falta {label}" for key, label in INVENTORY_HEADERS.items() if not fields[key]]
        rows.append(EquipmentRow(
            excel_row=excel_row, serial=fields["noseriecpu"], brand_model="", characteristics="",
            item=fields["items"], monitor_serial=fields["noseriemonitor"],
            responsible=fields["responsable"], inventory_format=True,
            status="Inválido" if errors else "Pendiente", error="; ".join(errors),
        ))
    if not rows:
        raise ValueError("El archivo no contiene filas de equipos.")
    return rows


def read_paired_rows(
    raw_headers: Sequence[object], iterator: Iterable[Sequence[object]]
) -> list[EquipmentRow] | None:
    headers = tuple(normalize_header(header) for header in raw_headers)
    if "cpu" not in headers and "monitores" not in headers:
        return None
    if headers[:6] != PAIRED_HEADERS:
        raise ValueError(
            "Las primeras seis columnas deben ser: CPU, Marca y Modelo, "
            "Características del equipo, Monitores, Marca y Modelo, "
            "Características del equipo."
        )
    rows: list[EquipmentRow] = []
    labels = (
        "CPU", "Marca y Modelo del CPU", "Características del CPU",
        "Monitores", "Marca y Modelo del monitor", "Características del monitor",
    )
    for excel_row, values in enumerate(iterator, start=2):
        fields = [clean_cell(values[index]) if index < len(values) else "" for index in range(6)]
        if not any(fields) and not any(clean_cell(value) for value in values[6:]):
            continue
        errors = [f"Falta {label}" for field, label in zip(fields, labels) if not field]
        rows.append(EquipmentRow(
            excel_row=excel_row, serial=fields[0], brand_model=fields[1],
            characteristics=fields[2], monitor_serial=fields[3],
            monitor_brand_model=fields[4], monitor_characteristics=fields[5],
            paired_format=True, status="Inválido" if errors else "Pendiente",
            error="; ".join(errors),
        ))
    if not rows:
        raise ValueError("El Excel no contiene filas de equipos.")
    return rows


def read_equipment_excel(path: Path) -> list[EquipmentRow]:
    if path.suffix.casefold() == ".csv":
        return read_ticket_csv(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration as exc:
            raise ValueError("El Excel está vacío.") from exc

        paired_rows = read_paired_rows(raw_headers, iterator)
        if paired_rows is not None:
            return paired_rows

        inventory_rows = read_inventory_rows(raw_headers, iterator)
        if inventory_rows is not None:
            return inventory_rows

        header_map: dict[str, int] = {}
        for index, header in enumerate(raw_headers):
            normalized = normalize_header(header)
            if normalized in LEGACY_HEADERS or normalized in TICKET_HEADERS or normalized in {"lab", "laboratorio"}:
                header_map[normalized] = index
        has_legacy_format = all(key in header_map for key in LEGACY_HEADERS)
        has_ticket_format = all(key in header_map for key in TICKET_HEADERS)
        if not has_legacy_format and not has_ticket_format:
            raise ValueError(
                "El Excel debe contener CPU, Marca y Modelo, Características del equipo, "
                "Monitores, Marca y Modelo, Características del equipo; "
                "o bien Items, No. serie CPU, No. serie monitor y Responsable; "
                "o bien NumeroSerie, Marcaymodelo y Caracteristicas; "
                "o bien Codigo, Marca, Modelo y Serie."
            )

        rows: list[EquipmentRow] = []
        seen_serials: set[str] = set()
        for excel_row, values in enumerate(iterator, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if has_ticket_format:
                code = clean_cell(values[header_map["codigo"]])
                brand = clean_cell(values[header_map["marca"]])
                model = clean_cell(values[header_map["modelo"]])
                serial = clean_cell(values[header_map["serie"]])
                lab_index = header_map.get("lab", header_map.get("laboratorio"))
                lab = clean_cell(values[lab_index]) if lab_index is not None else ""
                brand_model = f"{brand} {model}".strip()
                characteristics = ""
            else:
                code = brand = model = lab = ""
                serial = clean_cell(values[header_map["numeroserie"]])
                brand_model = clean_cell(values[header_map["marcaymodelo"]])
                characteristics = clean_cell(values[header_map["caracteristicas"]])
            errors: list[str] = []
            if has_ticket_format:
                if not code:
                    errors.append("Falta Codigo")
                if not brand:
                    errors.append("Falta Marca")
                if not model:
                    errors.append("Falta Modelo")
                if not serial:
                    errors.append("Falta Serie")
            else:
                if not serial:
                    errors.append("Falta NumeroSerie")
                if not brand_model:
                    errors.append("Falta Marcaymodelo")
                if not characteristics:
                    errors.append("Falta Caracteristicas")
            serial_key = serial.casefold()
            if serial and serial_key in seen_serials:
                errors.append("NumeroSerie duplicado dentro del Excel")
            if serial:
                seen_serials.add(serial_key)
            rows.append(
                EquipmentRow(
                    excel_row=excel_row,
                    serial=serial,
                    brand_model=brand_model,
                    characteristics=characteristics,
                    status="Inválido" if errors else "Pendiente",
                    error="; ".join(errors),
                    code=code,
                    brand=brand,
                    model=model,
                    lab=lab,
                )
            )
        if not rows:
            raise ValueError("El Excel no contiene filas de equipos.")
        return rows
    finally:
        workbook.close()


def read_ticket_csv(path: Path) -> list[EquipmentRow]:
    """Lee tickets o inventarios de CPU y monitores desde CSV."""
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(handle, dialect)
            raw_headers = next(reader, None)
            if raw_headers is None:
                raise ValueError("El archivo está vacío.")
            paired_rows = read_paired_rows(raw_headers, reader)
            if paired_rows is not None:
                return paired_rows
            inventory_rows = read_inventory_rows(raw_headers, reader)
            if inventory_rows is not None:
                return inventory_rows
            header_map = {normalize_header(header): index for index, header in enumerate(raw_headers)}
            if not all(key in header_map for key in TICKET_HEADERS):
                raise ValueError(
                    "El CSV debe contener Items, No. serie CPU, No. serie monitor y Responsable; "
                    "o bien Codigo, Marca, Modelo y Serie."
                )

            rows: list[EquipmentRow] = []
            seen_serials: set[str] = set()
            for excel_row, values in enumerate(reader, start=2):
                if not any(value.strip() for value in values):
                    continue
                def value_for(key: str) -> str:
                    index = header_map[key]
                    return values[index].strip() if index < len(values) else ""

                code, brand, model, serial = (value_for("codigo"), value_for("marca"), value_for("modelo"), value_for("serie"))
                lab = value_for("lab") if "lab" in header_map else value_for("laboratorio") if "laboratorio" in header_map else ""
                errors = [label for value, label in ((code, "Falta Codigo"), (brand, "Falta Marca"), (model, "Falta Modelo"), (serial, "Falta Serie")) if not value]
                if serial and serial.casefold() in seen_serials:
                    errors.append("Serie duplicada dentro del archivo")
                if serial:
                    seen_serials.add(serial.casefold())
                rows.append(EquipmentRow(excel_row, serial, f"{brand} {model}".strip(), "", status="Inválido" if errors else "Pendiente", error="; ".join(errors), code=code, brand=brand, model=model, lab=lab))
    except UnicodeDecodeError as exc:
        raise ValueError("El CSV debe estar guardado en UTF-8.") from exc
    if not rows:
        raise ValueError("El archivo no contiene filas de tickets.")
    return rows


def assign_verification_schedule(
    rows: list[EquipmentRow], verification_date: date, start: str, end: str
) -> None:
    valid_rows = [row for row in rows if row.valid]
    times = distribute_uniform_times(start, end, len(valid_rows))
    for row, assigned_time in zip(valid_rows, times):
        row.verification_date = verification_date
        row.verification_time = assigned_time


def build_subject(row: EquipmentRow) -> str:
    if row.paired_format:
        return f"Instalacion de CPU {row.serial} y monitor {row.monitor_serial} para su respectivo uso en el Laboratorio 1 de posgrado."
    if row.inventory_format:
        return f"Preparacion de CPU {row.serial} para su respectivo uso."
    if row.code and row.lab:
        return f"Informe Técnico PC {row.code} - {row.lab} Biblioteca RUSB"
    identifier = f"código {row.code} / serie {row.serial}" if row.code else f"número de serie del equipo {row.serial}"
    return f"Verificación del {identifier} entregado por COMTECH"


def build_plain_body(row: EquipmentRow) -> str:
    assert row.verification_date is not None
    if row.code and row.lab:
        return (
            "Por medio del presente se informa la realización del diagnóstico técnico "
            "del siguiente equipo de la Biblioteca RUSB:\n\n"
            f"Código UNI: {row.code}\n"
            f"Marca: {row.brand}\n"
            f"Modelo: {row.model}\n"
            f"Serie: {row.serial}\n"
            f"Ubicación: {row.lab}, Biblioteca RUSB\n"
            "Usuario responsable: wilmer.davila@uni.edu.ni\n"
            "Cola: Diagnóstico\n\n"
            "En general, el hardware actual presenta limitaciones significativas para "
            "las exigencias de las aplicaciones modernas, lo que provoca un bajo "
            "desempeño del equipo.\n\n"
            "Con base en los diagnósticos realizados, el equipo se declara en estado "
            "de obsolescencia.\n\n"
            "Saludos cordiales.\n"
        )
    if row.paired_format:
        equipment_details = (
            f"CPU: {row.serial}\n"
            f"Marca y modelo del CPU: {row.brand_model}\n"
            f"Características del CPU: {row.characteristics}\n"
            f"Monitor: {row.monitor_serial}\n"
            f"Marca y modelo del monitor: {row.monitor_brand_model}\n"
            f"Características del monitor: {row.monitor_characteristics}\n"
        )
        return (
            "Realizada la instalacion del equipo en el Laboratorio 1 de posgrado "
            "para su respectivo uso, con la siguiente información:\n\n"
            f"{equipment_details}\n"
            f"Fecha de instalación: {row.verification_date.strftime('%d/%m/%Y')}\n"
            f"Hora de instalación: {row.verification_time}\n"
        )
    elif row.inventory_format:
        equipment_details = (
            f"Ítem: {row.item}\n"
            f"No. serie CPU: {row.serial}\n"
            f"No. serie monitor: {row.monitor_serial}\n"
            f"Responsable: {row.responsible}\n"
        )
    elif row.code:
        equipment_details = (
            f"Código: {row.code}\n"
            f"Marca: {row.brand}\n"
            f"Modelo: {row.model}\n"
            f"Serie: {row.serial}\n"
        )
    else:
        equipment_details = (
            f"Características del equipo: {row.characteristics}\n"
            f"N/S: {row.serial}\n"
            f"Marca y Modelo: {row.brand_model}\n"
        )
    return (
        "Realizada la preparación del equipo recibido en DATACENTER FM800 para su respectivo uso, con la siguiente información:\n\n"
        "Proveedor de los equipos: COMTECH\n"
        "En colaboración con: Bodega RUSB\n"
        f"{equipment_details}\n"
        f"Fecha de preparación: {row.verification_date.strftime('%d/%m/%Y')}\n"
        f"Hora de preparación: {row.verification_time}\n"
    )


def build_html_body(row: EquipmentRow) -> str:
    assert row.verification_date is not None
    values = {
        "serial": html.escape(row.serial),
        "brand_model": html.escape(row.brand_model),
        "characteristics": html.escape(row.characteristics),
        "date": row.verification_date.strftime("%d/%m/%Y"),
        "time": html.escape(row.verification_time),
        "code": html.escape(row.code),
        "brand": html.escape(row.brand),
        "model": html.escape(row.model),
        "lab": html.escape(row.lab),
        "item": html.escape(row.item),
        "monitor_serial": html.escape(row.monitor_serial),
        "monitor_brand_model": html.escape(row.monitor_brand_model),
        "monitor_characteristics": html.escape(row.monitor_characteristics),
        "responsible": html.escape(row.responsible),
    }
    if row.code and row.lab:
        return f"""<!doctype html>
<html lang="es"><body style="margin:0;padding:24px;background:#f4f6f8;font-family:Arial,sans-serif;color:#1f2937;">
<div style="max-width:680px;margin:0 auto;background:#fff;border:1px solid #dbe2ea;border-radius:8px;overflow:hidden;">
<div style="background:#17365d;color:#fff;padding:18px 24px;"><h2 style="margin:0;font-size:20px;">Informe técnico de diagnóstico</h2></div>
<div style="padding:24px;"><p style="margin-top:0;">Estimado equipo de Diagnóstico:</p>
<p>Por medio del presente se informa la realización del diagnóstico técnico del siguiente equipo de la Biblioteca RUSB:</p>
<table role="presentation" style="width:100%;border-collapse:collapse;margin-top:18px;">
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;width:38%;">Código UNI</td><td style="padding:10px;border:1px solid #dbe2ea;">{values['code']}</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">Marca</td><td style="padding:10px;border:1px solid #dbe2ea;">{values['brand']}</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">Modelo</td><td style="padding:10px;border:1px solid #dbe2ea;">{values['model']}</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">Serie</td><td style="padding:10px;border:1px solid #dbe2ea;">{values['serial']}</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">Ubicación</td><td style="padding:10px;border:1px solid #dbe2ea;">{values['lab']}, Biblioteca RUSB</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">Usuario responsable</td><td style="padding:10px;border:1px solid #dbe2ea;">wilmer.davila@uni.edu.ni</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">Cola</td><td style="padding:10px;border:1px solid #dbe2ea;">Diagnóstico</td></tr>
</table>
<p>En general, el hardware actual presenta limitaciones significativas para las exigencias de las aplicaciones modernas, lo que provoca un bajo desempeño del equipo.</p>
<p>Con base en los diagnósticos realizados, el equipo se declara en estado de obsolescencia.</p>
<p style="margin-bottom:0;">Saludos cordiales.</p></div></div></body></html>"""
    if row.paired_format:
        equipment_rows = "\n".join(
            f'<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">{label}</td>'
            f'<td style="padding:10px;border:1px solid #dbe2ea;">{values[key]}</td></tr>'
            for label, key in (
                ("CPU", "serial"), ("Marca y modelo del CPU", "brand_model"),
                ("Características del CPU", "characteristics"),
                ("Monitor", "monitor_serial"),
                ("Marca y modelo del monitor", "monitor_brand_model"),
                ("Características del monitor", "monitor_characteristics"),
            )
        )
    elif row.inventory_format:
        equipment_rows = "\n".join(
            f'<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">{label}</td>'
            f'<td style="padding:10px;border:1px solid #dbe2ea;">{values[key]}</td></tr>'
            for label, key in (("Ítem", "item"), ("No. serie CPU", "serial"),
                               ("No. serie monitor", "monitor_serial"), ("Responsable", "responsible"))
        )
    elif row.code:
        equipment_rows = f"""<tr><td style=\"padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;width:38%;\">Código</td><td style=\"padding:10px;border:1px solid #dbe2ea;\">{values['code']}</td></tr>
<tr><td style=\"padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;\">Marca</td><td style=\"padding:10px;border:1px solid #dbe2ea;\">{values['brand']}</td></tr>
<tr><td style=\"padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;\">Modelo</td><td style=\"padding:10px;border:1px solid #dbe2ea;\">{values['model']}</td></tr>
<tr><td style=\"padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;\">Serie</td><td style=\"padding:10px;border:1px solid #dbe2ea;\">{values['serial']}</td></tr>"""
    else:
        equipment_rows = f"""<tr><td style=\"padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;width:38%;\">Características del equipo</td><td style=\"padding:10px;border:1px solid #dbe2ea;\">{values['characteristics']}</td></tr>
<tr><td style=\"padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;\">N/S</td><td style=\"padding:10px;border:1px solid #dbe2ea;\">{values['serial']}</td></tr>
<tr><td style=\"padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;\">Marca y Modelo</td><td style=\"padding:10px;border:1px solid #dbe2ea;\">{values['brand_model']}</td></tr>"""
    if row.paired_format:
        title = "Instalación de equipo"
        introduction = "Realizada la instalacion del equipo en el Laboratorio 1 de posgrado para su respectivo uso, con la siguiente información:"
        source_rows = ""
        date_label, time_label = "Fecha de instalación", "Hora de instalación"
    else:
        title = "Preparación de equipo"
        introduction = "Realizada la preparación del equipo recibido en DATACENTER FM800 para su respectivo uso, con la siguiente información:"
        source_rows = """<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;width:38%;">Proveedor de los equipos</td><td style="padding:10px;border:1px solid #dbe2ea;">COMTECH</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">En colaboración con</td><td style="padding:10px;border:1px solid #dbe2ea;">Bodega RUSB</td></tr>"""
        date_label, time_label = "Fecha de preparación", "Hora de preparación"
    return f"""<!doctype html>
<html lang="es"><body style="margin:0;padding:24px;background:#f4f6f8;font-family:Arial,sans-serif;color:#1f2937;">
<div style="max-width:680px;margin:0 auto;background:#fff;border:1px solid #dbe2ea;border-radius:8px;overflow:hidden;">
<div style="background:#17365d;color:#fff;padding:18px 24px;"><h2 style="margin:0;font-size:20px;">{title}</h2></div>
<div style="padding:24px;"><p style="margin-top:0;">{introduction}</p>
<table role="presentation" style="width:100%;border-collapse:collapse;margin-top:18px;">
{source_rows}
{equipment_rows}
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">{date_label}</td><td style="padding:10px;border:1px solid #dbe2ea;">{values['date']}</td></tr>
<tr><td style="padding:10px;border:1px solid #dbe2ea;background:#eef3f8;font-weight:bold;">{time_label}</td><td style="padding:10px;border:1px solid #dbe2ea;">{values['time']}</td></tr>
</table></div></div></body></html>"""


def graph_recipient(address: str) -> dict[str, dict[str, str]]:
    return {"emailAddress": {"address": address}}


def build_graph_payload(row: EquipmentRow, settings: MailSettings) -> dict[str, object]:
    return {
        "message": {
            "subject": build_subject(row),
            "body": {"contentType": "HTML", "content": build_html_body(row)},
            "toRecipients": [graph_recipient(settings.recipient)],
            "ccRecipients": [graph_recipient(address) for address in settings.cc],
            "bccRecipients": [graph_recipient(address) for address in settings.bcc],
        },
        "saveToSentItems": True,
    }


def load_oauth_config() -> dict[str, str]:
    if not OAUTH_CONFIG_FILE.exists():
        return {}
    try:
        data = json.loads(OAUTH_CONFIG_FILE.read_text(encoding="utf-8"))
        return {
            "tenant_id": str(data.get("tenant_id", "")).strip(),
            "client_id": str(data.get("client_id", "")).strip(),
        }
    except (OSError, ValueError, TypeError):
        return {}


def save_oauth_config(tenant_id: str, client_id: str) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    OAUTH_CONFIG_FILE.write_text(
        json.dumps({"tenant_id": tenant_id, "client_id": client_id}, indent=2),
        encoding="utf-8",
    )


class OAuthSession:
    """Sesión delegada de Microsoft mantenida únicamente en memoria."""

    def __init__(self, tenant_id: str, client_id: str) -> None:
        self.tenant_id = tenant_id.strip()
        self.client_id = client_id.strip()
        self.account_email = ""
        self.display_name = ""
        self._application: object | None = None
        self._account: dict[str, object] | None = None

    def sign_in(self) -> str:
        try:
            import msal
        except ImportError as exc:
            raise RuntimeError(
                "Falta la dependencia msal. Ejecute: python -m pip install -r requirements.txt"
            ) from exc

        authority = f"https://login.microsoftonline.com/{self.tenant_id}"
        application = msal.PublicClientApplication(self.client_id, authority=authority)
        result = application.acquire_token_interactive(
            scopes=GRAPH_SCOPES,
            prompt="select_account",
            timeout=300,
        )
        if "access_token" not in result:
            description = result.get("error_description") or result.get("error")
            raise RuntimeError(str(description or "Microsoft rechazó el inicio de sesión.").strip())

        claims = result.get("id_token_claims") or {}
        username = str(claims.get("preferred_username") or claims.get("email") or "").strip()
        accounts = application.get_accounts(username=username) if username else application.get_accounts()
        if not accounts:
            raise RuntimeError("Microsoft inició sesión, pero no devolvió una cuenta utilizable.")
        self._application = application
        self._account = accounts[0]
        self.account_email = str(self._account.get("username") or username).strip()
        self.display_name = str(claims.get("name") or self.account_email).strip()
        if not self.account_email:
            raise RuntimeError("No fue posible identificar el correo de la cuenta conectada.")
        return self.account_email

    def acquire_access_token(self) -> str:
        if self._application is None or self._account is None:
            raise RuntimeError("No hay una cuenta Microsoft conectada.")
        result = self._application.acquire_token_silent(GRAPH_SCOPES, account=self._account)
        if not result or "access_token" not in result:
            raise RuntimeError("La sesión de Microsoft expiró. Cierre sesión e ingrese nuevamente.")
        return str(result["access_token"])

    def sign_out(self) -> None:
        if self._application is not None and self._account is not None:
            self._application.remove_account(self._account)
        self._application = None
        self._account = None
        self.account_email = ""
        self.display_name = ""


class GraphMailClient:
    def __init__(self, oauth_session: OAuthSession, opener: Callable[..., object] = urlopen) -> None:
        self.oauth_session = oauth_session
        self.opener = opener

    def send(self, row: EquipmentRow, settings: MailSettings) -> None:
        token = self.oauth_session.acquire_access_token()
        payload = json.dumps(build_graph_payload(row, settings), ensure_ascii=False).encode("utf-8")
        request = Request(
            GRAPH_SEND_MAIL_URL,
            data=payload,
            headers={
                "Authorization": f"Bearer {token}",
                "Content-Type": "application/json; charset=utf-8",
            },
            method="POST",
        )
        try:
            with self.opener(request, timeout=30) as response:
                status = getattr(response, "status", None)
                if status is None:
                    status = response.getcode()
                if status != 202:
                    raise RuntimeError(f"Microsoft Graph respondió con estado inesperado {status}.")
        except HTTPError as exc:
            body = exc.read(8192).decode("utf-8", errors="replace")
            try:
                error = json.loads(body).get("error", {})
                detail = f"{error.get('code', 'Error')}: {error.get('message', body)}"
            except ValueError:
                detail = body or exc.reason
            raise RuntimeError(f"Microsoft Graph {exc.code}: {detail}") from exc
        except URLError as exc:
            raise RuntimeError(f"No fue posible conectar con Microsoft Graph: {exc.reason}") from exc


class HistoryStore:
    headers = [
        "FechaEnvioReal", "Modo", "Remitente", "Destinatario", "FilaExcel",
        "NumeroSerie", "MarcaYModelo", "FechaVerificacion", "HoraVerificacion",
        "Estado", "Error",
        "IdentificadorRegistro", "Items", "SerieMonitor", "Responsable",
        "CaracteristicasCPU", "MarcaYModeloMonitor", "CaracteristicasMonitor",
    ]

    def __init__(self, path: Path = HISTORY_FILE) -> None:
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def sent_keys(self) -> set[tuple[str, str, str, str]]:
        if not self.path.exists():
            return set()
        keys: set[tuple[str, str, str, str]] = set()
        with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
            for record in csv.DictReader(handle):
                if record.get("Estado") == "ENVIADO":
                    keys.add((
                        record.get("Remitente", "").casefold(),
                        record.get("Destinatario", "").casefold(),
                        record.get("IdentificadorRegistro") or record.get("NumeroSerie", "").casefold(),
                        record.get("FechaVerificacion", ""),
                    ))
        return keys

    def append(self, row: EquipmentRow, settings: MailSettings, state: str, error: str = "") -> None:
        is_new = not self.path.exists() or self.path.stat().st_size == 0
        fieldnames = self.headers
        if not is_new:
            with self.path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                fieldnames = list(reader.fieldnames or [])
                missing = [header for header in self.headers if header not in fieldnames]
                if missing:
                    records = list(reader)
            if missing:
                fieldnames += missing
                temporary_path = None
                try:
                    with tempfile.NamedTemporaryFile(
                        mode="w", encoding="utf-8-sig", newline="", dir=self.path.parent, delete=False,
                    ) as handle:
                        temporary_path = Path(handle.name)
                        writer = csv.DictWriter(handle, fieldnames=fieldnames)
                        writer.writeheader()
                        writer.writerows(records)
                    temporary_path.replace(self.path)
                finally:
                    if temporary_path is not None:
                        temporary_path.unlink(missing_ok=True)
        with self.path.open("a", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            if is_new:
                writer.writeheader()
            writer.writerow({
                "FechaEnvioReal": datetime.now().isoformat(timespec="seconds"),
                "Modo": "PRUEBA" if settings.test_mode else "REAL",
                "Remitente": settings.sender,
                "Destinatario": settings.recipient,
                "FilaExcel": row.excel_row,
                "NumeroSerie": row.serial,
                "MarcaYModelo": row.brand_model,
                "FechaVerificacion": row.verification_date.isoformat() if row.verification_date else "",
                "HoraVerificacion": row.verification_time,
                "Estado": state,
                "Error": error,
                "IdentificadorRegistro": row.history_identifier,
                "Items": row.item,
                "SerieMonitor": row.monitor_serial,
                "Responsable": row.responsible,
                "CaracteristicasCPU": row.characteristics,
                "MarcaYModeloMonitor": row.monitor_brand_model,
                "CaracteristicasMonitor": row.monitor_characteristics,
            })


class VerificationApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1180x790")
        self.minsize(1000, 700)
        self.rows: list[EquipmentRow] = []
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.cancel_event = threading.Event()
        self.worker: threading.Thread | None = None
        self.oauth_session: OAuthSession | None = None
        self._build_variables()
        self._configure_style()
        self._build_ui()
        self.after(150, self._drain_events)

    def _build_variables(self) -> None:
        oauth_config = load_oauth_config()
        self.file_var = tk.StringVar()
        self.tenant_var = tk.StringVar(value=oauth_config.get("tenant_id", ""))
        self.client_var = tk.StringVar(value=oauth_config.get("client_id", ""))
        self.account_var = tk.StringVar(value="Sin iniciar sesión")
        self.recipient_var = tk.StringVar()
        self.cc_var = tk.StringVar()
        self.bcc_var = tk.StringVar()
        self.date_var = tk.StringVar(value=datetime.now().date().isoformat())
        self.start_var = tk.StringVar(value="08:00")
        self.end_var = tk.StringVar(value="18:00")
        self.delay_var = tk.StringVar(value="5")
        self.test_var = tk.BooleanVar(value=True)
        self.summary_var = tk.StringVar(value="Seleccione un Excel y pulse Cargar y validar.")
        self.progress_var = tk.DoubleVar(value=0)
        self.progress_text_var = tk.StringVar(value="0 / 0")

    def _configure_style(self) -> None:
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Title.TLabel", font=("Arial", 17, "bold"), foreground="#17365d")
        style.configure("Section.TLabelframe", padding=10)
        style.configure("Section.TLabelframe.Label", font=("Arial", 10, "bold"), foreground="#17365d")
        style.configure("Primary.TButton", font=("Arial", 10, "bold"), padding=(12, 8))
        style.configure("Connected.TLabel", foreground="#176b3a", font=("Arial", 9, "bold"))
        style.configure("Treeview", rowheight=25)
        style.configure("Treeview.Heading", font=("Arial", 9, "bold"))

    def _build_ui(self) -> None:
        outer = ttk.Frame(self, padding=16)
        outer.pack(fill="both", expand=True)
        ttk.Label(outer, text=APP_TITLE, style="Title.TLabel").pack(anchor="w")
        ttk.Label(outer, text="La cuenta que inicia sesión en Microsoft será el remitente de los correos.").pack(anchor="w", pady=(2, 12))

        settings = ttk.Frame(outer)
        settings.pack(fill="x")
        settings.columnconfigure(0, weight=1)
        settings.columnconfigure(1, weight=1)
        source_frame = ttk.LabelFrame(settings, text="Archivo y verificación", style="Section.TLabelframe")
        source_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        source_frame.columnconfigure(1, weight=1)
        self._labeled_entry(source_frame, 0, "Archivo de equipos", self.file_var, button=("Buscar", self._browse))
        self._labeled_entry(source_frame, 1, "Fecha", self.date_var)
        self._labeled_entry(source_frame, 2, "Hora inicial", self.start_var)
        self._labeled_entry(source_frame, 3, "Hora final", self.end_var)
        self._labeled_entry(source_frame, 4, "Delay real (seg.)", self.delay_var)

        mail_frame = ttk.LabelFrame(settings, text="Microsoft 365 y destinatarios", style="Section.TLabelframe")
        mail_frame.grid(row=0, column=1, sticky="nsew", padx=(6, 0))
        mail_frame.columnconfigure(1, weight=1)
        self.tenant_entry = self._labeled_entry(mail_frame, 0, "Tenant ID", self.tenant_var)
        self.client_entry = self._labeled_entry(mail_frame, 1, "Client ID", self.client_var)
        account_line = ttk.Frame(mail_frame)
        account_line.grid(row=2, column=0, columnspan=3, sticky="ew", pady=(5, 7))
        account_line.columnconfigure(1, weight=1)
        ttk.Label(account_line, text="Cuenta:").grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Label(account_line, textvariable=self.account_var, style="Connected.TLabel").grid(row=0, column=1, sticky="w")
        self.login_button = ttk.Button(account_line, text="Iniciar sesión con Microsoft", command=self._sign_in)
        self.login_button.grid(row=0, column=2, padx=(8, 4))
        self.logout_button = ttk.Button(account_line, text="Cerrar sesión", command=self._sign_out, state="disabled")
        self.logout_button.grid(row=0, column=3)
        self._labeled_entry(mail_frame, 3, "Destinatario", self.recipient_var)
        self._labeled_entry(mail_frame, 4, "CC (opcional)", self.cc_var)
        self._labeled_entry(mail_frame, 5, "BCC (opcional)", self.bcc_var)

        actions = ttk.Frame(outer)
        actions.pack(fill="x", pady=12)
        ttk.Checkbutton(actions, text="Modo prueba (no inicia sesión ni envía)", variable=self.test_var, command=self._refresh_send_button).pack(side="left")
        ttk.Button(actions, text="Cargar y validar", command=self._load, style="Primary.TButton").pack(side="right")
        self.send_button = ttk.Button(actions, command=self._start, style="Primary.TButton", state="disabled")
        self.send_button.pack(side="right", padx=8)
        self.preview_button = ttk.Button(actions, text="Vista previa", command=self._preview, state="disabled")
        self.preview_button.pack(side="right")
        self.cancel_button = ttk.Button(actions, text="Cancelar", command=self._cancel, state="disabled")
        self.cancel_button.pack(side="right", padx=8)
        self._refresh_send_button()

        ttk.Label(outer, textvariable=self.summary_var).pack(anchor="w", pady=(0, 6))
        progress_line = ttk.Frame(outer)
        progress_line.pack(fill="x", pady=(0, 8))
        ttk.Progressbar(progress_line, variable=self.progress_var, maximum=100).pack(side="left", fill="x", expand=True)
        ttk.Label(progress_line, textvariable=self.progress_text_var, width=12, anchor="e").pack(side="right", padx=(8, 0))

        notebook = ttk.Notebook(outer)
        notebook.pack(fill="both", expand=True)
        table_tab, log_tab = ttk.Frame(notebook), ttk.Frame(notebook)
        notebook.add(table_tab, text="Equipos")
        notebook.add(log_tab, text="Actividad")
        columns = ("order", "row", "code", "brand", "model", "lab", "serial", "time", "status", "error")
        self.tree = ttk.Treeview(table_tab, columns=columns, show="headings")
        headings = {"order": "N.º", "row": "Fila Excel", "code": "Código", "brand": "Marca", "model": "Modelo", "lab": "Ubicación", "serial": "Serie", "time": "Hora", "status": "Estado", "error": "Detalle"}
        widths = {"order": 55, "row": 75, "code": 145, "brand": 110, "model": 130, "lab": 95, "serial": 120, "time": 105, "status": 100, "error": 260}
        for column in columns:
            self.tree.heading(column, text=headings[column])
            self.tree.column(column, width=widths[column], anchor="w" if column in ("brand", "model", "lab", "error") else "center")
        table_scroll = ttk.Scrollbar(table_tab, orient="vertical", command=self.tree.yview)
        table_scroll_x = ttk.Scrollbar(table_tab, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=table_scroll.set, xscrollcommand=table_scroll_x.set)
        table_scroll_x.pack(side="bottom", fill="x")
        self.tree.pack(side="left", fill="both", expand=True)
        table_scroll.pack(side="right", fill="y")
        self.log_text = tk.Text(log_tab, wrap="word", state="disabled", font=("TkFixedFont", 9))
        log_scroll = ttk.Scrollbar(log_tab, orient="vertical", command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=log_scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        log_scroll.pack(side="right", fill="y")

    def _labeled_entry(self, parent: ttk.Widget, row: int, label: str, variable: tk.StringVar, button: tuple[str, Callable[[], None]] | None = None) -> ttk.Entry:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=3)
        entry = ttk.Entry(parent, textvariable=variable)
        entry.grid(row=row, column=1, sticky="ew", pady=3)
        if button:
            ttk.Button(parent, text=button[0], command=button[1]).grid(row=row, column=2, padx=(6, 0), pady=3)
        return entry

    def _browse(self) -> None:
        selected = filedialog.askopenfilename(title="Seleccionar archivo de equipos", filetypes=(("Excel o CSV", "*.xlsx *.xlsm *.csv"), ("Todos los archivos", "*.*")))
        if selected:
            self.file_var.set(selected)

    def _validated_schedule_inputs(self) -> tuple[Path, date, str, str]:
        path = Path(self.file_var.get().strip()).expanduser()
        if not path.is_file():
            raise ValueError("Seleccione un archivo de equipos existente.")
        selected_date = parse_date(self.date_var.get())
        parse_clock(self.start_var.get())
        parse_clock(self.end_var.get())
        return path, selected_date, self.start_var.get().strip(), self.end_var.get().strip()

    def _sign_in(self) -> None:
        if self.worker and self.worker.is_alive():
            return
        tenant_id, client_id = self.tenant_var.get().strip(), self.client_var.get().strip()
        if not tenant_id or not client_id:
            messagebox.showerror("Faltan datos OAuth", "Introduzca el Tenant ID y Client ID proporcionados por el administrador.", parent=self)
            return
        try:
            save_oauth_config(tenant_id, client_id)
        except OSError as exc:
            self._log(f"No se pudo guardar la configuración OAuth: {exc}")
        self.login_button.configure(state="disabled")
        self.account_var.set("Esperando inicio de sesión en el navegador...")
        self._log("Abriendo el inicio de sesión oficial de Microsoft en el navegador.")
        self.worker = threading.Thread(target=self._sign_in_worker, args=(tenant_id, client_id), daemon=True)
        self.worker.start()

    def _sign_in_worker(self, tenant_id: str, client_id: str) -> None:
        session = OAuthSession(tenant_id, client_id)
        try:
            session.sign_in()
            self.events.put(("login_success", session))
        except Exception as exc:
            self.events.put(("login_error", str(exc)))

    def _sign_out(self) -> None:
        if self.oauth_session is not None:
            self.oauth_session.sign_out()
        self.oauth_session = None
        self.account_var.set("Sin iniciar sesión")
        self.tenant_entry.configure(state="normal")
        self.client_entry.configure(state="normal")
        self.login_button.configure(state="normal")
        self.logout_button.configure(state="disabled")
        self._log("Sesión de Microsoft cerrada en la aplicación.")
        self._refresh_send_button()

    def _load(self) -> None:
        try:
            path, selected_date, start, end = self._validated_schedule_inputs()
            rows = read_equipment_excel(path)
            assign_verification_schedule(rows, selected_date, start, end)
        except Exception as exc:
            messagebox.showerror("No se pudo cargar", str(exc), parent=self)
            return
        self.rows = rows
        self._refresh_table()
        valid = sum(row.valid for row in rows)
        invalid = len(rows) - valid
        self.summary_var.set(f"Registros: {len(rows)} | Válidos: {valid} | Inválidos: {invalid}")
        self.preview_button.configure(state="normal" if valid else "disabled")
        self.progress_var.set(0)
        self.progress_text_var.set(f"0 / {valid}")
        self._log(f"Excel cargado: {path.name}. {valid} filas válidas y {invalid} inválidas.")
        self._refresh_send_button()

    def _refresh_table(self) -> None:
        self.tree.delete(*self.tree.get_children())
        paired = any(row.paired_format for row in self.rows)
        inventory = any(row.inventory_format for row in self.rows)
        if paired:
            columns = ("order", "row", "serial", "brand_model", "characteristics", "monitor_serial", "monitor_brand_model", "monitor_characteristics", "time", "status", "error")
            headings = ("N.º", "Fila Excel", "CPU", "Marca y modelo CPU", "Características CPU", "Monitor", "Marca y modelo monitor", "Características monitor", "Hora", "Estado", "Detalle")
            widths = (55, 75, 130, 210, 225, 130, 210, 230, 105, 100, 260)
        elif inventory:
            columns = ("order", "row", "item", "serial", "monitor_serial", "responsible", "time", "status", "error")
            headings = ("N.º", "Fila Excel", "Ítem", "No. serie CPU", "No. serie monitor", "Responsable", "Hora", "Estado", "Detalle")
            widths = (55, 75, 55, 135, 145, 110, 105, 100, 260)
        else:
            columns = ("order", "row", "code", "brand", "model", "lab", "serial", "time", "status", "error")
            headings = ("N.º", "Fila Excel", "Código", "Marca", "Modelo", "Ubicación", "Serie", "Hora", "Estado", "Detalle")
            widths = (55, 75, 145, 110, 130, 95, 120, 105, 100, 260)
        self.tree.configure(columns=columns)
        for column, heading, width in zip(columns, headings, widths):
            self.tree.heading(column, text=heading)
            self.tree.column(column, width=width, anchor="w" if column in ("brand", "model", "lab", "responsible", "brand_model", "monitor_brand_model", "characteristics", "monitor_characteristics", "error") else "center")
        for order, row in enumerate(self.rows, start=1):
            brand = row.brand or row.brand_model
            if paired:
                values = (order, row.excel_row, row.serial, row.brand_model, row.characteristics, row.monitor_serial, row.monitor_brand_model, row.monitor_characteristics, row.verification_time, row.status, row.error)
            elif inventory:
                values = (order, row.excel_row, row.item, row.serial, row.monitor_serial, row.responsible, row.verification_time, row.status, row.error)
            else:
                values = (order, row.excel_row, row.code, brand, row.model, row.lab, row.serial, row.verification_time, row.status, row.error)
            self.tree.insert("", "end", iid=str(order - 1), values=values)

    def _preview(self) -> None:
        row = next((item for item in self.rows if item.valid), None)
        if not row:
            messagebox.showinfo("Vista previa", "No hay filas válidas.", parent=self)
            return
        window = tk.Toplevel(self)
        window.title("Vista previa del primer correo")
        window.geometry("760x620")
        container = ttk.Frame(window, padding=16)
        container.pack(fill="both", expand=True)
        ttk.Label(container, text="Asunto", font=("Arial", 10, "bold")).pack(anchor="w")
        subject = tk.Text(container, height=2, wrap="word")
        subject.insert("1.0", build_subject(row)); subject.configure(state="disabled"); subject.pack(fill="x", pady=(4, 12))
        ttk.Label(container, text="Contenido", font=("Arial", 10, "bold")).pack(anchor="w")
        body = tk.Text(container, wrap="word")
        body.insert("1.0", build_plain_body(row)); body.configure(state="disabled"); body.pack(fill="both", expand=True, pady=(4, 10))
        ttk.Button(container, text="Cerrar", command=window.destroy).pack(anchor="e")

    def _mail_settings(self) -> MailSettings:
        recipient = parse_address_list(self.recipient_var.get(), "Destinatario", allow_empty=False)
        if len(recipient) != 1:
            raise ValueError("Destinatario debe contener una sola dirección.")
        cc = parse_address_list(self.cc_var.get(), "CC")
        bcc = parse_address_list(self.bcc_var.get(), "BCC")
        try:
            delay = float(self.delay_var.get().replace(",", "."))
        except ValueError as exc:
            raise ValueError("El delay debe ser un número.") from exc
        if delay < 0:
            raise ValueError("El delay no puede ser negativo.")
        if self.test_var.get():
            sender = self.oauth_session.account_email if self.oauth_session else "modo-prueba@local"
        else:
            if self.oauth_session is None or not self.oauth_session.account_email:
                raise ValueError("Inicie sesión con Microsoft antes de realizar un envío real.")
            sender = self.oauth_session.account_email
        return MailSettings(sender, recipient[0], cc, bcc, delay, self.test_var.get())

    def _start(self) -> None:
        if self.worker and self.worker.is_alive():
            return
        try:
            _, selected_date, start, end = self._validated_schedule_inputs()
            assign_verification_schedule(self.rows, selected_date, start, end)
            settings = self._mail_settings()
        except Exception as exc:
            messagebox.showerror("Configuración inválida", str(exc), parent=self)
            return
        self._refresh_table()
        valid_rows = [row for row in self.rows if row.valid]
        if not valid_rows:
            messagebox.showinfo("Sin registros", "No hay filas válidas para procesar.", parent=self)
            return
        mode = "SIMULACIÓN" if settings.test_mode else "ENVÍO REAL MEDIANTE MICROSOFT GRAPH"
        detail = (f"Modo: {mode}\nRegistros válidos: {len(valid_rows)}\nRemitente autenticado: {settings.sender}\nDestino: {settings.recipient}\nCC: {', '.join(settings.cc) or '(vacío)'}\nBCC: {', '.join(settings.bcc) or '(vacío)'}\n\n¿Desea continuar?")
        if not messagebox.askyesno("Confirmar proceso", detail, parent=self):
            return
        self.cancel_event.clear()
        self.send_button.configure(state="disabled")
        self.preview_button.configure(state="disabled")
        self.login_button.configure(state="disabled")
        self.logout_button.configure(state="disabled")
        self.cancel_button.configure(state="normal")
        self.worker = threading.Thread(target=self._process_rows, args=(valid_rows, settings, self.oauth_session), daemon=True)
        self.worker.start()

    def _process_rows(self, rows: list[EquipmentRow], settings: MailSettings, oauth_session: OAuthSession | None) -> None:
        history = HistoryStore()
        sent_keys = history.sent_keys()
        graph = GraphMailClient(oauth_session) if oauth_session else None
        total, completed, fatal_error = len(rows), 0, ""
        try:
            if not settings.test_mode:
                if graph is None:
                    fatal_error = "No hay una sesión de Microsoft conectada."
                    return
                try:
                    oauth_session.acquire_access_token()
                    self.events.put(("log", f"Token válido para {settings.sender}."))
                except Exception as exc:
                    fatal_error = f"No fue posible usar la sesión de Microsoft: {exc}"
                    self.events.put(("log", fatal_error))
                    return
            for row in self.rows:
                if not row.valid:
                    history.append(row, settings, "INVALIDO", row.error)
            for index, row in enumerate(rows, start=1):
                if self.cancel_event.is_set():
                    self.events.put(("log", "Proceso cancelado por el usuario.")); break
                assert row.verification_date is not None
                duplicate_key = (settings.sender.casefold(), settings.recipient.casefold(), row.history_identifier, row.verification_date.isoformat())
                if duplicate_key in sent_keys:
                    row.status, row.error = "Omitido", "Ya fue enviado anteriormente"
                    history.append(row, settings, "OMITIDO", row.error)
                    self.events.put(("row", (self.rows.index(row), row.status, row.error)))
                    self.events.put(("log", f"{row.serial}: omitido por duplicado."))
                    completed += 1; self.events.put(("progress", (completed, total))); continue
                try:
                    if settings.test_mode:
                        row.status = "Simulado"; history.append(row, settings, "SIMULADO")
                    else:
                        assert graph is not None
                        graph.send(row, settings)
                        row.status = "Enviado"; sent_keys.add(duplicate_key); history.append(row, settings, "ENVIADO")
                    row.error = ""
                    self.events.put(("log", f"{row.serial}: {row.status.lower()} correctamente."))
                except Exception as exc:
                    row.status, row.error = "Error", str(exc)
                    history.append(row, settings, "ERROR", row.error)
                    self.events.put(("log", f"{row.serial}: ERROR - {row.error}"))
                completed += 1
                self.events.put(("row", (self.rows.index(row), row.status, row.error)))
                self.events.put(("progress", (completed, total)))
                if not settings.test_mode and index < total and not self.cancel_event.is_set() and self.cancel_event.wait(settings.delay_seconds):
                    self.events.put(("log", "Proceso cancelado durante la espera.")); break
        finally:
            self.events.put(("finished", (completed, total, settings.test_mode, self.cancel_event.is_set(), fatal_error)))

    def _cancel(self) -> None:
        if self.worker and self.worker.is_alive():
            self.cancel_event.set(); self.cancel_button.configure(state="disabled")
            self._log("Cancelación solicitada; se detendrá después de la operación actual.")

    def _drain_events(self) -> None:
        try:
            while True:
                event, payload = self.events.get_nowait()
                if event == "log": self._log(str(payload))
                elif event == "login_success":
                    assert isinstance(payload, OAuthSession)
                    self.worker = None
                    self.oauth_session = payload; self.account_var.set(payload.account_email)
                    self.tenant_entry.configure(state="disabled"); self.client_entry.configure(state="disabled")
                    self.login_button.configure(state="disabled"); self.logout_button.configure(state="normal")
                    self._log(f"Sesión iniciada correctamente como {payload.account_email}."); self._refresh_send_button()
                elif event == "login_error":
                    self.worker = None
                    self.oauth_session = None; self.account_var.set("Sin iniciar sesión")
                    self.login_button.configure(state="normal"); self.logout_button.configure(state="disabled")
                    self._refresh_send_button(); messagebox.showerror("No se pudo iniciar sesión", str(payload), parent=self)
                elif event == "progress":
                    completed, total = payload  # type: ignore[misc]
                    self.progress_var.set((completed / total * 100) if total else 0); self.progress_text_var.set(f"{completed} / {total}")
                elif event == "row":
                    row_index, status, error = payload  # type: ignore[misc]
                    item = str(row_index)
                    if self.tree.exists(item):
                        values = list(self.tree.item(item, "values")); values[-2], values[-1] = status, error; self.tree.item(item, values=values)
                elif event == "finished":
                    completed, total, test_mode, cancelled, fatal_error = payload  # type: ignore[misc]
                    self.worker = None
                    self.cancel_button.configure(state="disabled"); self.preview_button.configure(state="normal")
                    if self.oauth_session:
                        self.logout_button.configure(state="normal"); self.login_button.configure(state="disabled")
                    else: self.login_button.configure(state="normal")
                    self._refresh_send_button()
                    if fatal_error: title = "No se inició el envío"; self.summary_var.set(str(fatal_error))
                    elif cancelled: title = "Proceso cancelado"; self.summary_var.set(f"{title}: {completed} de {total} registros procesados.")
                    else: title = "Simulación finalizada" if test_mode else "Envío finalizado"; self.summary_var.set(f"{title}: {completed} de {total} registros procesados.")
                    messagebox.showinfo(title, self.summary_var.get(), parent=self)
        except queue.Empty:
            pass
        self.after(150, self._drain_events)

    def _log(self, message: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"[{datetime.now().strftime('%H:%M:%S')}] {message}\n")
        self.log_text.see("end"); self.log_text.configure(state="disabled")

    def _refresh_send_button(self) -> None:
        if not hasattr(self, "send_button"): return
        self.send_button.configure(text="Simular envío" if self.test_var.get() else "Iniciar envío real")
        can_send = any(row.valid for row in self.rows) and (self.test_var.get() or self.oauth_session is not None)
        if self.worker and self.worker.is_alive(): can_send = False
        self.send_button.configure(state="normal" if can_send else "disabled")


def main() -> None:
    VerificationApp().mainloop()


if __name__ == "__main__":
    main()
