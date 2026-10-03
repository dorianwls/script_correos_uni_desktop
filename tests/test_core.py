import tempfile
import csv
import queue
import threading
import sys
import types
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from verificador_comtech.app import (
    EquipmentRow,
    GraphMailClient,
    HistoryStore,
    MailSettings,
    OAuthSession,
    VerificationApp,
    assign_verification_schedule,
    build_graph_payload,
    build_subject,
    build_plain_body,
    distribute_uniform_times,
    parse_address_list,
    read_equipment_excel,
)


class CoreTests(unittest.TestCase):
    def test_paired_cpu_monitor_excel_reads_repeated_headers_and_email_details(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "equipos.xlsx"
            wb = Workbook()
            wb.active.append(["CPU", "Marca y Modelo", "Características del equipo", "Monitores", "Marca y Modelo", "Características del equipo"])
            wb.active.append(["MZ03WGLX", "Lenovo ThinkCentre M70s Gen 6", "Ultra 7 265 / 32 GB / 1 TB / W11P", "V90G0NDP", "Lenovo ThinkVision T22i-30", "22 pulgadas; DisplayPort / HDMI / VGA"])
            wb.active.append(["MZ03WGKC", "Lenovo ThinkCentre M70s Gen 6", "Ultra 7 265 / 32 GB / 1 TB / W11P", "V90G0NPR", "Lenovo ThinkVision T22i-30", "22 pulgadas; DisplayPort / HDMI / VGA"])
            wb.save(path)
            rows = read_equipment_excel(path)
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row.valid and row.paired_format for row in rows))
        self.assertEqual((rows[0].serial, rows[0].monitor_serial), ("MZ03WGLX", "V90G0NDP"))
        self.assertEqual(rows[0].monitor_brand_model, "Lenovo ThinkVision T22i-30")
        self.assertEqual(rows[0].monitor_characteristics, "22 pulgadas; DisplayPort / HDMI / VGA")
        self.assertNotEqual(rows[0].history_identifier, rows[1].history_identifier)
        assign_verification_schedule(rows, date(2026, 10, 2), "08:00", "18:00")
        plain = build_plain_body(rows[0])
        html_body = build_graph_payload(rows[0], MailSettings("a@example.com", "b@example.com", (), (), 0, True))["message"]["body"]["content"]
        for detail in ("MZ03WGLX", "V90G0NDP", "Lenovo ThinkCentre M70s Gen 6", "Lenovo ThinkVision T22i-30", "Ultra 7 265 / 32 GB / 1 TB / W11P", "22 pulgadas; DisplayPort / HDMI / VGA"):
            self.assertIn(detail, plain)
            self.assertIn(detail, html_body)
        self.assertEqual(
            build_subject(rows[0]),
            "Instalacion de CPU MZ03WGLX y monitor V90G0NDP para su respectivo uso en el Laboratorio 1 de posgrado.",
        )
        for body in (plain, html_body):
            self.assertIn("Realizada la instalacion del equipo en el Laboratorio 1 de posgrado para su respectivo uso", body)
            self.assertIn("Fecha de instalación", body)
            self.assertIn("Hora de instalación", body)
            self.assertNotIn("Proveedor de los equipos", body)
            self.assertNotIn("En colaboración con", body)
            self.assertNotIn("preparación", body.casefold())

    def test_paired_excel_marks_missing_cells_and_rejects_wrong_headers(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "equipos.xlsx"
            wb = Workbook()
            wb.active.append(["CPU", "Marca y Modelo", "Características del equipo", "Monitores", "Marca y Modelo", "Características del equipo"])
            wb.active.append(["CPU001", "Marca PC", "i7", "MON001", "Marca monitor", None])
            wb.save(path)
            rows = read_equipment_excel(path)
            self.assertFalse(rows[0].valid)
            self.assertIn("Características del monitor", rows[0].error)
            wb.active.cell(1, 4, "Monitor")
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "primeras seis columnas"):
                read_equipment_excel(path)

    def test_paired_csv_and_html_escape(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "equipos.csv"
            path.write_text(
                "CPU;Marca y Modelo;Características del equipo;Monitores;Marca y Modelo;Características del equipo\n"
                "0001;Lenovo & Dell;i7;0002;Pantalla <22>;HDMI\n",
                encoding="utf-8-sig",
            )
            row = read_equipment_excel(path)[0]
        self.assertTrue(row.valid)
        self.assertEqual((row.serial, row.monitor_serial), ("0001", "0002"))
        assign_verification_schedule([row], date(2026, 10, 2), "08:00", "18:00")
        body = build_graph_payload(row, MailSettings("a@example.com", "b@example.com", (), (), 0, True))["message"]["body"]["content"]
        self.assertIn("Lenovo &amp; Dell", body)
        self.assertIn("Pantalla &lt;22&gt;", body)

    def test_responsible_workbooks_keep_all_40_rows_and_unique_emails(self):
        folder = Path(__file__).resolve().parents[1] / "equipos_por_responsable"
        all_rows = []
        settings = MailSettings("sender@example.com", "recipient@example.com", (), (), 0, True)
        for name, count, first_item in (("Rollman", 12, 1), ("Suyen", 14, 13), ("Dorian", 14, 27)):
            rows = read_equipment_excel(folder / f"equipos_{name.lower()}.xlsx")
            self.assertEqual(len(rows), count)
            self.assertTrue(all(row.valid for row in rows))
            self.assertEqual([row.item for row in rows], [str(i) for i in range(first_item, first_item + count)])
            self.assertEqual({row.responsible for row in rows}, {name})
            assign_verification_schedule(rows, date(2026, 9, 29), "08:00", "18:00")
            for row in rows:
                payload = build_graph_payload(row, settings)
                for value in (row.serial, row.monitor_serial, row.responsible):
                    self.assertIn(value, payload["message"]["body"]["content"])
                    self.assertIn(value, build_plain_body(row))
                self.assertEqual(payload["message"]["subject"], f"Preparacion de CPU {row.serial} para su respectivo uso.")
                self.assertIn(f"Ítem: {row.item}", build_plain_body(row))
                self.assertIn("Preparación de equipo", payload["message"]["body"]["content"])
                for body in (build_plain_body(row), payload["message"]["body"]["content"]):
                    self.assertIn("Realizada la preparación del equipo recibido en DATACENTER FM800 para su respectivo uso", body)
                    self.assertIn("Fecha de preparación", body)
                    self.assertIn("Hora de preparación", body)
                    self.assertNotIn("verificación", body.casefold())
            all_rows.extend(rows)
        self.assertEqual(len(all_rows), 40)
        self.assertEqual(len({row.history_identifier for row in all_rows}), 40)
        self.assertEqual(all_rows[34].serial, all_rows[37].serial)
        self.assertEqual(all_rows[31].serial, all_rows[38].serial)
        self.assertEqual(all_rows[30].monitor_serial, all_rows[38].monitor_serial)

    def test_inventory_missing_fields_and_csv_header_normalization(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "inventory.csv"
            path.write_text(
                "Items;No. serie CPU;No. serie monitor;Responsable\n"
                "1;CPU001;MON001;Rollman\n"
                "2;CPU001;MON002;Suyen\n"
                "3;CPU003;;Dorian\n"
                "4;CPU004;MON004;\n"
                ";CPU005;MON005;Dorian\n"
                "6;;MON006;Dorian\n"
                ";;;\n", encoding="utf-8-sig",
            )
            rows = read_equipment_excel(path)
            self.assertEqual(len(rows), 6)
            self.assertTrue(rows[0].valid)
            self.assertTrue(rows[1].valid)
            for row, missing in zip(rows[2:], ("No. serie monitor", "Responsable", "Items", "No. serie CPU")):
                self.assertFalse(row.valid)
                self.assertIn(missing, row.error)
            assign_verification_schedule(rows, date(2026, 9, 29), "08:00", "18:00")
            self.assertEqual(rows[0].verification_time, "08:00 AM")
            self.assertEqual(rows[1].verification_time, "01:30 PM")
            path.write_text("Items,No. serie CPU,Responsable\n1,CPU001,Rollman\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "No. serie monitor"):
                read_equipment_excel(path)

    def test_inventory_email_escapes_html_and_preserves_serial_zeroes(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "inventory.xlsx"
            wb = Workbook()
            wb.active.append(["Responsable", "No. serie monitor", "Items", "No. serie CPU"])
            wb.active.append(["A & B <Test>", "00002", 1, "00001"])
            wb.save(path)
            row = read_equipment_excel(path)[0]
        assign_verification_schedule([row], date(2026, 9, 29), "08:00", "18:00")
        payload = build_graph_payload(row, MailSettings("a@example.com", "b@example.com", (), (), 0, True))
        body = payload["message"]["body"]["content"]
        self.assertIn("A &amp; B &lt;Test&gt;", body)
        self.assertIn("00001", body)
        self.assertIn("00002", body)

    def test_inventory_history_preserves_old_records_and_distinguishes_rows_and_senders(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "history.csv"
            old_headers = HistoryStore.headers[:11]
            with path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=old_headers)
                writer.writeheader()
                writer.writerow({"Remitente": "a@example.com", "Destinatario": "b@example.com", "NumeroSerie": "CPU001", "FechaVerificacion": "2026-09-29", "Estado": "ENVIADO"})
            history = HistoryStore(path)
            legacy_key = ("a@example.com", "b@example.com", "cpu001", "2026-09-29")
            self.assertIn(legacy_key, history.sent_keys())
            rows = [EquipmentRow(2 + i, "CPU001", "", "", item=str(i + 1), monitor_serial="MON001", responsible="Dorian", inventory_format=True) for i in range(2)]
            assign_verification_schedule(rows, date(2026, 9, 29), "08:00", "18:00")
            settings = MailSettings("a@example.com", "b@example.com", (), (), 0, False)
            history.append(rows[0], settings, "ENVIADO")
            history.append(rows[1], settings, "SIMULADO")
            key = lambda sender, row: (sender, settings.recipient, row.history_identifier, "2026-09-29")
            self.assertIn(legacy_key, history.sent_keys())
            self.assertIn(key(settings.sender, rows[0]), history.sent_keys())
            self.assertNotIn(key(settings.sender, rows[1]), history.sent_keys())
            self.assertNotIn(key("other@example.com", rows[0]), history.sent_keys())
            with path.open(encoding="utf-8-sig", newline="") as handle:
                records = list(csv.DictReader(handle))
            self.assertEqual(len(records), 3)
            self.assertEqual(records[0]["NumeroSerie"], "CPU001")
            self.assertEqual(records[1]["SerieMonitor"], "MON001")
            self.assertEqual(records[1]["Responsable"], "Dorian")

    def test_inventory_batch_sends_every_row_then_blocks_repeat_for_same_sender(self):
        rows = read_equipment_excel(Path(__file__).resolve().parents[1] / "equipos_por_responsable/equipos_dorian.xlsx")
        assign_verification_schedule(rows, date(2026, 9, 29), "08:00", "18:00")
        fake_app = types.SimpleNamespace(rows=rows, events=queue.Queue(), cancel_event=threading.Event())
        fake_session = types.SimpleNamespace(acquire_access_token=lambda: "fake-token")
        with tempfile.TemporaryDirectory() as temp:
            history = HistoryStore(Path(temp) / "history.csv")
            with patch("verificador_comtech.app.HistoryStore", return_value=history), patch("verificador_comtech.app.GraphMailClient") as graph:
                settings = MailSettings("dorian@example.com", "recipient@example.com", (), (), 0, False)
                VerificationApp._process_rows(fake_app, rows, settings, fake_session)
                self.assertEqual(graph.return_value.send.call_count, 14)
                self.assertTrue(all(row.status == "Enviado" for row in rows))
                VerificationApp._process_rows(fake_app, rows, settings, fake_session)
                self.assertEqual(graph.return_value.send.call_count, 14)
                self.assertTrue(all(row.status == "Omitido" for row in rows))
                for row in rows:
                    row.error = ""
                other_settings = MailSettings("other@example.com", "recipient@example.com", (), (), 0, False)
                VerificationApp._process_rows(fake_app, rows, other_settings, fake_session)
                self.assertEqual(graph.return_value.send.call_count, 28)

    def test_600_uniform_times_exclude_lunch(self):
        values = distribute_uniform_times("08:00", "18:00", 600)
        self.assertEqual(len(values), 600)
        self.assertEqual(len(set(values)), 600)
        self.assertEqual(values[0], "08:00:00 AM")
        self.assertEqual(values[-1], "05:59:06 PM")
        parsed = [datetime.strptime(value, "%I:%M:%S %p") for value in values]
        self.assertEqual(parsed, sorted(parsed))
        self.assertTrue(all(value.hour != 12 for value in parsed))
        self.assertEqual(parsed[267].strftime("%H:%M:%S"), "13:00:18")

    def test_lunch_boundaries_and_partial_overlap(self):
        for start, end, expected in (
            ("11:00", "14:00", ["11:00 AM", "11:30 AM", "01:00 PM", "01:30 PM"]),
            ("12:00", "14:00", ["01:00 PM", "01:15 PM", "01:30 PM", "01:45 PM"]),
            ("12:30", "14:00", ["01:00 PM", "01:15 PM", "01:30 PM", "01:45 PM"]),
            ("11:00", "12:30", ["11:00 AM", "11:15 AM", "11:30 AM", "11:45 AM"]),
            ("11:00", "12:00", ["11:00 AM", "11:15 AM", "11:30 AM", "11:45 AM"]),
            ("13:00", "14:00", ["01:00 PM", "01:15 PM", "01:30 PM", "01:45 PM"]),
        ):
            with self.subTest(start=start, end=end):
                self.assertEqual(distribute_uniform_times(start, end, 4), expected)
        for start, end in (("12:00", "13:00"), ("12:15", "12:45")):
            with self.subTest(start=start, end=end):
                with self.assertRaisesRegex(ValueError, "almuerzo"):
                    distribute_uniform_times(start, end, 1)
        with self.assertRaisesRegex(ValueError, "fuera del almuerzo"):
            distribute_uniform_times("11:59", "13:01", 121)

    def test_seconds_are_used_for_short_range(self):
        values = distribute_uniform_times("08:00", "08:01", 60)
        self.assertEqual(values[0], "08:00:00 AM")
        self.assertEqual(values[-1], "08:00:59 AM")

    def test_invalid_range(self):
        with self.assertRaises(ValueError):
            distribute_uniform_times("18:00", "08:00", 10)

    def test_addresses(self):
        parsed = parse_address_list("a@example.com; b@example.com", "CC")
        self.assertEqual(parsed, ("a@example.com", "b@example.com"))

    def test_read_and_schedule(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "equipos.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["NumeroSerie", "Marcaymodelo", "Caracteristicas"])
            ws.append(["ABC001", "Dell Latitude", "i5 / 16 GB"])
            ws.append(["ABC002", "HP ProBook", "i7 / 32 GB"])
            wb.save(path)
            rows = read_equipment_excel(path)
            assign_verification_schedule(rows, date(2026, 9, 22), "08:00", "18:00")
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0].verification_time, "08:00 AM")
            self.assertEqual(rows[1].verification_time, "01:30 PM")
            self.assertIn("ABC001", build_subject(rows[0]))

    def test_read_ticket_format_and_build_ticket_body(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "tickets.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Codigo", "Marca", "Modelo", "Serie"])
            ws.append(["07-1217-7001-83", "Dell", "XPS", "D7BLXM2"])
            wb.save(path)
            rows = read_equipment_excel(path)

        assign_verification_schedule(rows, date(2026, 9, 22), "08:00", "18:00")
        self.assertEqual(rows[0].code, "07-1217-7001-83")
        self.assertEqual(rows[0].brand, "Dell")
        self.assertEqual(rows[0].model, "XPS")
        self.assertEqual(rows[0].serial, "D7BLXM2")
        self.assertIn("código 07-1217-7001-83", build_subject(rows[0]))
        payload = build_graph_payload(
            rows[0], MailSettings("a@empresa.com", "b@empresa.com", (), (), 0, True)
        )
        body = payload["message"]["body"]["content"]
        self.assertIn("Código</td>", body)
        self.assertIn("07-1217-7001-83", body)
        self.assertIn("Modelo</td>", body)

    def test_graph_payload_uses_recipients_and_html(self):
        row = EquipmentRow(
            excel_row=2,
            serial="ABC<001>",
            brand_model="Dell Latitude",
            characteristics="i5 & 16 GB",
            verification_date=date(2026, 9, 22),
            verification_time="08:00 AM",
        )
        settings = MailSettings(
            sender="usuario@empresa.com",
            recipient="destino@example.com",
            cc=("copia@example.com",),
            bcc=("oculta@example.com",),
            delay_seconds=5,
            test_mode=False,
        )
        payload = build_graph_payload(row, settings)
        message = payload["message"]
        self.assertEqual(message["toRecipients"][0]["emailAddress"]["address"], "destino@example.com")
        self.assertEqual(message["ccRecipients"][0]["emailAddress"]["address"], "copia@example.com")
        self.assertEqual(message["bccRecipients"][0]["emailAddress"]["address"], "oculta@example.com")
        self.assertIn("ABC&lt;001&gt;", message["body"]["content"])
        self.assertIn("i5 &amp; 16 GB", message["body"]["content"])
        self.assertTrue(payload["saveToSentItems"])

    def test_oauth_interactive_session_and_silent_token(self):
        class FakeApplication:
            def __init__(self, client_id, authority):
                self.client_id = client_id
                self.authority = authority
                self.removed = False

            def acquire_token_interactive(self, **kwargs):
                return {
                    "access_token": "initial-token",
                    "id_token_claims": {
                        "preferred_username": "usuario@uni.edu.ni",
                        "name": "Usuario UNI",
                    },
                }

            def get_accounts(self, username=None):
                return [{"username": username or "usuario@uni.edu.ni"}]

            def acquire_token_silent(self, scopes, account):
                return {"access_token": "silent-token"}

            def remove_account(self, account):
                self.removed = True

        fake_msal = types.SimpleNamespace(PublicClientApplication=FakeApplication)
        with patch.dict(sys.modules, {"msal": fake_msal}):
            session = OAuthSession("tenant-id", "client-id")
            self.assertEqual(session.sign_in(), "usuario@uni.edu.ni")
            self.assertEqual(session.acquire_access_token(), "silent-token")
            session.sign_out()
            self.assertEqual(session.account_email, "")

    def test_graph_client_posts_to_me_sendmail(self):
        row = EquipmentRow(
            excel_row=2,
            serial="ABC001",
            brand_model="Dell Latitude",
            characteristics="i5 / 16 GB",
            verification_date=date(2026, 9, 22),
            verification_time="08:00 AM",
        )
        settings = MailSettings(
            "usuario@uni.edu.ni",
            "destino@example.com",
            (),
            (),
            5,
            False,
        )
        captured = {}

        class FakeSession:
            def acquire_access_token(self):
                return "test-token"

        class FakeResponse:
            status = 202

            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, traceback):
                return False

        def fake_opener(request, timeout):
            captured["url"] = request.full_url
            captured["authorization"] = request.get_header("Authorization")
            captured["body"] = request.data.decode("utf-8")
            return FakeResponse()

        GraphMailClient(FakeSession(), opener=fake_opener).send(row, settings)
        self.assertEqual(captured["url"], "https://graph.microsoft.com/v1.0/me/sendMail")
        self.assertEqual(captured["authorization"], "Bearer test-token")
        self.assertIn("destino@example.com", captured["body"])


if __name__ == "__main__":
    unittest.main()
